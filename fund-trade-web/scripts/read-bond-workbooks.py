"""Read the two source bond workbooks through a captured private pipe only.

--read emits private JSON to its caller; never run it with inherited stdout.
--refresh emits counts only and atomically updates the existing encrypted file.
No plaintext payload, account value, password or data key is written to disk.
"""
from __future__ import annotations

import argparse
import base64
import csv
from contextlib import redirect_stdout, redirect_stderr
from datetime import date, datetime, time, timedelta, timezone
import hashlib
import importlib.util
import io
import json
import math
from pathlib import Path
import secrets
import sys

import openpyxl
from cryptography.hazmat.primitives.ciphers.aead import AESGCM

sys.dont_write_bytecode = True
CODES = ("006485", "006663")
SUFFIX = "债基交易账本.xlsx"
HEADERS = (
    "申请日期", "申请时间", "交易类型", "申请金额（元）", "申请份额", "确认份额",
    "净值日期", "确认净值", "实际手续费（元）", "确认转出金额（元）", "份额确认日期",
    "状态", "来源状态", "备注", "记录ID", "来源图片", "预计完成日期",
)
TYPE_MAP = {"买入": "buy", "定投": "sip", "卖出": "sell", "转换转入": "transfer-in", "转换转出": "transfer-out"}
STATUS_MAP = {"已确认": "confirmed", "待确认": "pending", "进行中": "pending", "撤销": "cancelled", "已撤销": "cancelled", "撤单": "cancelled"}
SOURCE_STATUSES = {"confirmed", "unverified", "processing", "cancelled"}


def fail():
    raise ValueError("Bond workbook source unavailable.")


def numeric(value, *, optional=False, positive=False):
    if value is None or value == "":
        if optional:
            return None
        fail()
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        fail()
    if positive and value <= 0:
        fail()
    return value


def iso_day(value, *, optional=False):
    if value is None or value == "":
        if optional:
            return None
        fail()
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, str):
        try:
            result = date.fromisoformat(value.strip())
            if result.isoformat() == value.strip():
                return result.isoformat()
        except ValueError:
            pass
    fail()


def iso_time(value, application_day):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        # XLSX serial dates can round an exact source second to .999 ms below
        # it. The source screenshots have second precision, so recover that
        # precision before comparing the application date or emitting a time.
        rounded = value + timedelta(microseconds=500000)
        if rounded.date().isoformat() != application_day:
            fail()
        return rounded.strftime("%H:%M:%S")
    if isinstance(value, time):
        rounded = datetime.combine(date.fromisoformat(application_day), value) + timedelta(microseconds=500000)
        return rounded.strftime("%H:%M:%S")
    if isinstance(value, str):
        try:
            parsed = datetime.strptime(value.strip(), "%H:%M:%S")
            if parsed.strftime("%H:%M:%S") == value.strip():
                return value.strip()
        except ValueError:
            pass
    fail()


def source_paths(workspace):
    result = {}
    for code in CODES:
        directories = [item for item in (workspace / "基金").glob(f"*/{code}_*") if item.is_dir()]
        if len(directories) != 1:
            fail()
        source = directories[0] / f"{code}_{SUFFIX}"
        if not source.is_file():
            fail()
        result[code] = source
    return result


def read_transactions(sheet, code):
    if tuple(cell.value for cell in sheet[5][:len(HEADERS)]) != HEADERS:
        fail()
    transactions, ids = [], set()
    for row in sheet.iter_rows(min_row=6, max_col=len(HEADERS), values_only=True):
        if all(value is None or value == "" for value in row):
            continue
        (application_date, application_time, raw_type, amount, requested_shares, shares,
         nav_date, nav, fee, confirmed_amount, confirmed_date, raw_status, source_status,
         note, identifier, _source_filename, completion_deadline) = row
        kind = TYPE_MAP.get(raw_type, raw_type)
        status = STATUS_MAP.get(raw_status, raw_status)
        if kind not in TYPE_MAP.values() or status not in STATUS_MAP.values():
            fail()
        if status == "cancelled":
            continue
        application_day = iso_day(application_date)
        application_clock = iso_time(application_time, application_day)
        submitted = f"{application_day} {application_clock}" if application_clock else None
        if not isinstance(identifier, str) or not identifier.strip() or len(identifier) > 200 or identifier in ids:
            fail()
        ids.add(identifier)
        if source_status not in SOURCE_STATUSES or (status == "confirmed" and source_status != "confirmed"):
            fail()
        numbers = {
            "amount": numeric(amount, optional=True), "requestedShares": numeric(requested_shares, optional=True),
            "shares": numeric(shares, optional=True), "nav": numeric(nav, optional=True),
            "fee": numeric(fee, optional=True), "confirmedAmount": numeric(confirmed_amount, optional=True),
        }
        if any(value is not None and value < 0 for value in numbers.values()):
            fail()
        net_value_date = iso_day(nav_date, optional=True)
        actual_confirmation_date = iso_day(confirmed_date, optional=True)
        deadline = iso_day(completion_deadline, optional=True)
        if status == "confirmed" and (numbers["shares"] is None or numbers["shares"] <= 0
                                      or numbers["nav"] is None or numbers["nav"] <= 0
                                      or numbers["fee"] is None or net_value_date is None):
            fail()
        if note is not None and not isinstance(note, str):
            fail()
        source_note = (note or "").strip()
        if deadline:
            source_note += f" 预计完成日期{deadline}，并非份额确认日期。"
        transactions.append({
            "id": identifier, "submittedAt": submitted, "date": application_day, "type": kind,
            **numbers, "navDate": net_value_date, "confirmedDate": actual_confirmation_date,
            "status": status, "sourceStatus": source_status, "note": source_note.strip(),
        })
    return sorted(transactions, key=lambda row: (row["date"], row["submittedAt"] or "", row["id"]))


def read_book(source, code):
    book = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        if not {"持仓总览", "逐笔交易"}.issubset(book.sheetnames):
            fail()
        overview = book["持仓总览"]
        if str(overview["C4"].value).zfill(6) != code:
            fail()
        market_value = numeric(overview["C7"].value)
        holding_profit = numeric(overview["C8"].value)
        if market_value < 0 or market_value - holding_profit <= 0:
            fail()
        timestamp = overview["C6"].value
        if not isinstance(timestamp, str):
            fail()
        # A leading apostrophe keeps the original ISO source identifier as text
        # instead of letting spreadsheet import normalize its timezone/precision.
        timestamp = timestamp.removeprefix("'")
        try:
            datetime.fromisoformat(timestamp.replace("Z", "+00:00"))
        except ValueError:
            fail()
        complete = overview["C13"].value
        if not isinstance(complete, bool):
            fail()
        lots = []
        for row in overview.iter_rows(min_row=17, max_col=3, values_only=True):
            if all(value is None or value == "" for value in row):
                continue
            confirmed_date, remaining_shares, status = row
            if status in ("pending", "cancelled", "待确认", "已撤销", "进行中"):
                continue
            if status not in ("confirmed", "已确认"):
                fail()
            remaining = numeric(remaining_shares)
            if remaining < 0:
                fail()
            if remaining > 0:
                lots.append({"confirmedDate": iso_day(confirmed_date), "remainingShares": remaining, "status": "confirmed"})
        return {"bond": {"code": code, "marketValue": market_value, "holdingProfit": holding_profit},
                "updatedAt": timestamp, "holdingLots": {"complete": complete, "lots": lots},
                "transactions": read_transactions(book["逐笔交易"], code)}
    finally:
        book.close()


def envelope_tools(root):
    spec = importlib.util.spec_from_file_location("bond_crypto_support", root / "scripts" / "encrypt-workbooks.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def read_original(root):
    module = envelope_tools(root)
    source = root / "data" / "bonds.enc.json"
    envelope = module.read_envelope(source)
    if envelope["version"] == 2:
        key = module.unwrap_daily_key(envelope, module.beijing_today_code())
    else:
        key = hashlib.pbkdf2_hmac("sha256", module.beijing_today_code().encode(), module.decode_b64(envelope["salt"]), module.ITERATIONS, dklen=32)
    encoded = AESGCM(key).decrypt(module.decode_b64(envelope["iv"]), module.decode_b64(envelope["ciphertext"]), None)
    result = json.loads(encoded.decode("utf-8"))
    if not isinstance(result, dict) or not isinstance(result.get("bonds"), list):
        fail()
    return result, envelope, key


def signature(sources):
    return [(str(source), source.stat().st_size, source.stat().st_mtime_ns) for source in sources]


def public_nav_path(source, code):
    return source.parent / "历史净值" / f"{code}_债基持有期计算净值.csv"


def read_public_navs(source, code):
    navs = {}
    with public_nav_path(source, code).open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if not {"基金代码", "净值日期", "单位净值", "数据来源"}.issubset(reader.fieldnames or []):
            fail()
        for row in reader:
            if row["基金代码"] != code:
                fail()
            day = iso_day(row["净值日期"])
            try:
                nav = float(row["单位净值"])
            except (ValueError, TypeError):
                fail()
            numeric(nav, positive=True)
            if day in navs or not row["数据来源"]:
                fail()
            navs[day] = nav
    if not navs:
        fail()
    return navs


def calculate_book(book, source, code, root, *, as_of=None):
    """Keep reported cells intact; only missing completed economics are derived."""
    spec = importlib.util.spec_from_file_location("bond_workbook_math", root / "scripts" / "bond-workbook-math.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    navs = read_public_navs(source, code)
    today = as_of or datetime.now(timezone(timedelta(hours=8))).date().isoformat()
    try:
        calculated = module.reconstruct(code, book["transactions"], navs, as_of=today)
    except module.ReconstructionError:
        # Never turn a short history or missing NAV into a full allocation.
        return book
    result = {**book, "transactions": calculated["transactions"]}
    if not book["holdingLots"]["complete"] and not book["holdingLots"]["lots"]:
        result["holdingLots"] = calculated["holdingLots"]
    return result


def read(workspace, root):
    sources = source_paths(workspace)
    dependencies = [*sources.values(), *(public_nav_path(source, code) for code, source in sources.items()),
                    root / "data" / "bonds.enc.json", root / "scripts" / "bond-workbook-math.py"]
    before = signature(dependencies)
    original, envelope, key = read_original(root)
    books = {code: calculate_book(read_book(source, code), source, code, root) for code, source in sources.items()}
    payload = {**original, "funds": {}, "bonds": [books[code]["bond"] for code in CODES],
               "updatedAt": max(books[code]["updatedAt"] for code in CODES),
               "holdingLots": {code: books[code]["holdingLots"] for code in CODES},
               "transactions": {code: books[code]["transactions"] for code in CODES}}
    if before != signature(dependencies):
        fail()
    return payload, envelope, key


def refresh(root, payload, envelope, key):
    # Reuse all daily wrappers and the authenticated data key. Existing browser
    # ciphertext stays decryptable after a source refresh or process restart.
    iv = secrets.token_bytes(12)
    plaintext = json.dumps(payload, ensure_ascii=False, separators=(",", ":"), allow_nan=False).encode("utf-8")
    ciphertext = AESGCM(key).encrypt(iv, plaintext, None)
    if AESGCM(key).decrypt(iv, ciphertext, None) != plaintext:
        fail()
    next_envelope = {**envelope, "iv": base64.b64encode(iv).decode(), "ciphertext": base64.b64encode(ciphertext).decode()}
    output = root / "data" / "bonds.enc.json"
    staging = root.parent / "临时文件" / root.name / "account-staging"
    staging.mkdir(parents=True, exist_ok=True)
    temporary = staging / (output.name + "." + secrets.token_hex(8) + ".tmp")
    temporary.write_text(json.dumps(next_envelope, separators=(",", ":")), encoding="utf-8")
    temporary.replace(output)


def main():
    parser = argparse.ArgumentParser(description="Read private bond workbook sources through a captured pipe.")
    parser.add_argument("--workspace", required=True, type=Path)
    parser.add_argument("--root", required=True, type=Path)
    action = parser.add_mutually_exclusive_group(required=True)
    action.add_argument("--read", action="store_true")
    action.add_argument("--refresh", action="store_true")
    args = parser.parse_args()
    with redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
        payload, envelope, key = read(args.workspace.resolve(), args.root.resolve())
        if args.refresh:
            refresh(args.root.resolve(), payload, envelope, key)
        encoded = json.dumps(payload, ensure_ascii=False, separators=(",", ":"), allow_nan=False)
    if args.read:
        sys.stdout.write(encoded)
    else:
        sys.stdout.write(json.dumps({"codes": list(CODES), "workbookSources": True,
                                    "transactionCounts": {code: len(payload["transactions"][code]) for code in CODES},
                                    "lotEvidenceComplete": {code: payload["holdingLots"][code]["complete"] for code in CODES},
                                    "stableEncryptionKeyPreserved": True}, separators=(",", ":")))


if __name__ == "__main__":
    try:
        main()
    except Exception:
        sys.stderr.write("Bond workbook read failed; no account values emitted.\n")
        sys.exit(1)
