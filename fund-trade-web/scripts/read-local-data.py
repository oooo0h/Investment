"""Read local sources in memory; stdout is a private pipe consumed by live-data.mjs.

Never invoke this reader with inherited stdout, and never log its JSON payload.
No workbook, source file, decrypted ledger or encryption key is written here.
"""
from __future__ import annotations

import argparse
from contextlib import redirect_stderr, redirect_stdout
import csv
from datetime import date, datetime, timedelta
import importlib.util
import io
import json
import math
from pathlib import Path
import re
import sys

sys.dont_write_bytecode = True
TEMPLATE_SUFFIX = "单笔交易自动计算_双模式_模板2026-09-28.xlsx"


class SourcePath(type(Path())):
    """The existing workbook reader must ignore Excel's temporary owner files."""

    def glob(self, pattern, **kwargs):
        return (item for item in super().glob(pattern, **kwargs) if not item.name.startswith("~$"))


def fail():
    raise ValueError("Local source is unavailable.")


def valid_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        for pattern in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
            try:
                return datetime.strptime(value.strip(), pattern).date().isoformat()
            except ValueError:
                pass
    if isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value):
        if value == int(value) and 1 <= value <= 2958465 and value != 60:
            return (date(1899, 12, 30) + timedelta(days=int(value) + (1 if value < 60 else 0))).isoformat()
    fail()


def public_number(value):
    if isinstance(value, bool) or value is None:
        fail()
    result = float(value)
    if not math.isfinite(result) or result <= 0:
        fail()
    return result


def fund_directory(workspace, code):
    matches = [item for item in (workspace / "基金").glob(f"*/{code}*") if item.is_dir()]
    if len(matches) != 1:
        fail()
    return matches[0]


def select_nav(directory, code):
    current = directory / f"{code}_全部历史净值.xlsx"
    if current.is_file():
        return current
    pattern = re.compile(rf"^{code}_全部历史净值_截至(\d{{4}}-\d{{2}}-\d{{2}})\.(csv|json|xlsx)$")
    priority = {"csv": 2, "json": 1, "xlsx": 0}
    candidates = []
    legacy = directory / "历史净值"
    for item in legacy.iterdir() if legacy.is_dir() else []:
        match = pattern.fullmatch(item.name)
        if match and item.is_file():
            try:
                cutoff = valid_date(match[1])
            except ValueError:
                continue
            candidates.append((cutoff, priority[match[2]], item))
    if not candidates:
        fail()
    return max(candidates, key=lambda item: item[:2])[2]


def tabular_rows(rows, code):
    iterator = iter(rows)
    header = next(iterator, None)
    if header is None:
        fail()
    names = [str(cell).strip() if cell is not None else "" for cell in header]
    aliases = {
        "date": ("净值日期", "日期", "date"),
        "nav": ("单位净值", "净值", "nav"),
        "code": ("基金代码", "code"),
        "source": ("数据来源", "source"),
    }
    indexes = {}
    for key, labels in aliases.items():
        matches = [index for index, name in enumerate(names) if name in labels]
        if len(matches) > 1 or (key in ("date", "nav") and len(matches) != 1):
            fail()
        if matches:
            indexes[key] = matches[0]
    history, source = [], "本地公开净值表"
    for values in iterator:
        if not values or all(value is None or str(value).strip() == "" for value in values):
            continue
        if len(values) != len(names):
            fail()
        if "code" in indexes:
            raw_code = values[indexes["code"]]
            if raw_code is None or str(raw_code).strip().zfill(6) != code:
                fail()
        history.append({"date": valid_date(values[indexes["date"]]), "nav": public_number(values[indexes["nav"]])})
        if "source" in indexes and values[indexes["source"]] not in (None, "") and source == "本地公开净值表":
            source = str(values[indexes["source"]]).strip()
    return history, source


def read_nav(source, code):
    if source.suffix.lower() == ".json":
        document = json.loads(source.read_text(encoding="utf-8-sig"))
        if not isinstance(document, dict) or str(document.get("code", code)).zfill(6) != code or not isinstance(document.get("rows"), list):
            fail()
        history = [{"date": valid_date(row.get("date")), "nav": public_number(row.get("nav"))} for row in document["rows"]]
        origin = document.get("source") if isinstance(document.get("source"), str) else "本地公开净值表"
    elif source.suffix.lower() == ".csv":
        with source.open("r", encoding="utf-8-sig", newline="") as handle:
            history, origin = tabular_rows(csv.reader(handle, strict=True), code)
    else:
        import openpyxl

        workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
        try:
            sheet = workbook[code] if code in workbook.sheetnames else workbook.worksheets[0]
            history, origin = tabular_rows(sheet.iter_rows(values_only=True), code)
        finally:
            workbook.close()
    if not history or any(index and row["date"] <= history[index - 1]["date"] for index, row in enumerate(history)):
        fail()
    return {"code": code, "source": origin, "asOf": history[-1]["date"], "history": history}


def signature(paths):
    return [(str(item), item.stat().st_size, item.stat().st_mtime_ns, item.stat().st_ino) for item in paths]


def read(workspace, root, codes):
    module_path = root / "scripts" / "encrypt-workbooks.py"
    envelope_path = root / "data" / "ledger.enc.json"
    selected, workbook_paths = {}, {}
    for code in codes:
        directory = fund_directory(workspace, code)
        selected[code] = select_nav(directory, code)
        candidates = [item for item in directory.iterdir() if item.is_file() and not item.name.startswith("~$") and item.name.endswith(TEMPLATE_SUFFIX)]
        if len(candidates) > 1:
            fail()
        if candidates:
            workbook_paths[code] = candidates[0]
    sources = [module_path, envelope_path, *selected.values(), *workbook_paths.values()]
    before = signature(sources)
    spec = importlib.util.spec_from_file_location("local_workbook_reader", module_path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    original = module.decrypt_payload(envelope_path, module.beijing_today_code())
    original_funds = original["funds"]
    # NAV coverage is broader than recorded account transactions. Missing records
    # mean an empty ledger, never an inferred purchase or fabricated holding.
    funds = {code: original_funds.get(code, {"transactions": [], "mode": "profit-first"}) for code in codes}
    for code in workbook_paths:
        funds[code] = {**funds[code], **module.read_fund(SourcePath(workspace), code), "incomplete": False}
    nav = {code: read_nav(selected[code], code) for code in codes}
    if before != signature(sources):
        fail()
    summary = [{
        "code": code,
        "nav": {"path": str(selected[code]), "asOf": nav[code]["asOf"], "recordCount": len(nav[code]["history"])},
        "transactions": {
            "kind": "workbook" if code in workbook_paths else "encrypted" if funds[code]["transactions"] else "empty",
            "path": str(workbook_paths[code]) if code in workbook_paths else str(envelope_path) if funds[code]["transactions"] else None,
            "recordCount": len(funds[code]["transactions"]),
        },
    } for code in codes]
    return {"vault": {"demo": False, "createdAt": original.get("createdAt"), "funds": funds}, "nav": nav, "sourceSummary": summary}


def main():
    parser = argparse.ArgumentParser(description="Read local sources through a private stdout pipe.")
    parser.add_argument("--workspace", required=True, type=Path)
    parser.add_argument("--root", required=True, type=Path)
    parser.add_argument("--codes", required=True)
    args = parser.parse_args()
    codes = tuple(args.codes.split(","))
    if not codes or len(set(codes)) != len(codes) or any(re.fullmatch(r"\d{6}", code) is None for code in codes):
        fail()
    # Existing helpers and dependency warnings must never leak to the terminal.
    with redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
        result = read(args.workspace.resolve(), args.root.resolve(), codes)
        encoded = json.dumps(result, ensure_ascii=False, separators=(",", ":"), allow_nan=False)
    sys.stdout.write(encoded)


if __name__ == "__main__":
    try:
        main()
    except Exception:
        sys.stderr.write("Local data read failed.\n")
        sys.exit(1)
