"""Migrate only public NAV/report folders after full saved-workbook verification."""

import sys
sys.dont_write_bytecode = True
from pathlib import Path
from decimal import Decimal
from datetime import datetime, date, timedelta, timezone
import csv
import hashlib
import json
import re
import shutil
import sys
import zipfile
import openpyxl
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
STATE = ROOT / "临时文件/工具脚本/单份净值整理_20261020"
ARCHIVE = ROOT / '公共资料/历史净值'
INDEX = ARCHIVE / '全基金历史净值当前入口.json'
CENTRAL = ARCHIVE / '维护数据'
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
rel = lambda p: p.relative_to(ROOT).as_posix()

def load(p):
    return json.loads(p.read_text(encoding='utf-8-sig'))

def save(p, d):
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(d, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')

def verify(doc, workbook):
    assert doc['count'] == len(doc['rows'])
    dates = [r['date'] for r in doc['rows']]
    assert dates == sorted(set(dates))
    assert (dates[0], dates[-1]) == (doc['first_date'], doc['last_date'])
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    sources = [('全部历史净值', doc['rows'])] + ([('待官网复核', doc['supplementary_rows'])] if doc.get('supplementary_rows') else [])
    assert all(name in wb.sheetnames for name, _ in sources)
    for name, source in sources:
        rows = list(wb[name].values)
        assert rows[0] == ('日期', '净值', '日增长率')
        assert len(rows) == len(source) + 1
        for actual, row in zip(rows[1:], source):
            day = actual[0].date().isoformat() if isinstance(actual[0], datetime) else (date(1899, 12, 30) + timedelta(days=int(actual[0]))).isoformat()
            assert day == row['date']
            assert abs(Decimal(str(actual[1])) - Decimal(str(row['nav']))) < Decimal('1e-12')
            if row.get('daily_growth_pct') in (None, ''):
                assert actual[2] is None
            else:
                assert abs(Decimal(str(actual[2])) * 100 - Decimal(str(row['daily_growth_pct']))) < Decimal('1e-10')
    wb.close()

def plan():
    idx = load(INDEX)
    items, pdfs = [], []
    bycode = {e['code']:e for e in idx['funds']}
    folders = sorted(p for group in (ROOT/'基金').iterdir() if group.is_dir() for p in group.iterdir() if p.is_dir() and re.match(r'^\d{6}_',p.name))
    assert len(folders) == len({p.name[:6] for p in folders})
    for folder in folders:
        code = folder.name[:6]
        history = folder/'历史净值'
        entry = bycode.get(code)
        master = ROOT/entry['json_path'] if entry else STATE/code/f'{code}_全部历史净值.json'
        doc = load(master)
        assert doc['code'] == code
        xlsx = STATE/f'{code}_全部历史净值.xlsx'
        if not xlsx.exists():
            # Temporarily include the newly downloaded fund for the exporter.
            assert code == '100055'
            entry = dict(code=code, name=doc['name'], records=doc['count'], first_date=doc['first_date'], last_date=doc['last_date'], source=doc['source'], json_path=rel(master), csv_path=rel(master.with_suffix('.csv')), json_sha256=sha(master))
            bycode[code] = entry
            idx['funds'] = sorted(bycode.values(),key=lambda e:e['code'])
            save(INDEX,idx)
            print('EXPORT_REQUIRED',code)
            return
        verify(doc,xlsx)
        originals = {rel(p):sha(p) for p in history.rglob('*') if p.is_file()} if history.exists() else {}
        for old in history.glob(code+'_全部历史净值*.json') if history.exists() else []:
            olddoc=load(old)
            assert {r['date'] for r in olddoc.get('rows',[])}.issubset({r['date'] for r in doc['rows']})
        items.append(dict(code=code, folder=rel(folder), history=rel(history), master=rel(master), staged=rel(xlsx), staged_sha256=sha(xlsx), originals=originals))
        reportdir=folder/'PDF报告'
        seen={}
        reports=sorted(reportdir.glob('*.pdf'),key=lambda p: ('official_' in p.name or p.name[0].isascii() and not re.match(r'^20\d{2}-',p.name), '摘要' in p.name, p.name)) if reportdir.exists() else []
        for p in reports:
            digest=sha(p)
            if digest in seen:
                pdfs.append(dict(source=rel(p), sha256=digest, duplicate_of=seen[digest]))
                continue
            match=re.search(r'(20\d{2})\s*年',p.name) or re.search(r'(?:report_|prospectus_)(20\d{2})',p.name) or re.search(r'(20\d{2})-\d{2}-\d{2}',p.name)
            if not match:
                text=' '.join(page.extract_text() or '' for page in PdfReader(p).pages[:2])
                match=re.search(r'(20\d{2})\s*年',text)
            assert match, f'Unresolved PDF year: {p.name}'
            year=match[1]
            filename=p.name
            if filename=='official_prospectus_2026-05-29.pdf': filename='2026-05-29_华夏全球科技先锋_更新招募说明书.pdf'
            if filename=='official_annual_report_2025.pdf': filename='2026-03-31_华夏全球科技先锋_2025年年度报告.pdf'
            target=reportdir/year/filename
            assert not target.exists()
            pdfs.append(dict(source=rel(p), target=rel(target), year=year, sha256=digest))
            seen[digest]=rel(target)
    save(STATE/'plan.json',dict(funds=items,pdfs=pdfs))
    print(json.dumps(dict(funds=len(items),nav_files_before=sum(len(f['originals']) for f in items),pdf_before=len(pdfs),pdf_retained=sum('target' in p for p in pdfs),pdf_duplicates=sum('duplicate_of' in p for p in pdfs))))

def apply():
    p=load(STATE/'plan.json')
    # Global preflight precedes all moves/deletions.
    for f in p['funds']:
        assert sha(ROOT/f['staged'])==f['staged_sha256']
        history=(ROOT/f['history']).resolve()
        assert history.is_relative_to((ROOT/'基金').resolve())
        assert history.name=='历史净值' and history.parent.name.startswith(f['code']+'_')
        current={rel(x):sha(x) for x in history.rglob('*') if x.is_file()} if history.exists() else {}
        assert current==f['originals'], 'Concurrent NAV change; replan first'
        verify(load(ROOT/f['master']),ROOT/f['staged'])
    for report in p['pdfs']: assert sha(ROOT/report['source'])==report['sha256']
    migration=[]
    for f in p['funds']:
        code=f['code']; old=ROOT/f['history']; central=CENTRAL/code
        assert not central.exists()
        central.mkdir(parents=True)
        master=ROOT/f['master']; doc=load(master)
        jp=central/f'{code}_全部历史净值.json'; cp=jp.with_suffix('.csv')
        shutil.copy2(master,jp)
        with cp.open('w',encoding='utf-8-sig',newline='') as handle:
            writer=csv.writer(handle,lineterminator='\r\n'); writer.writerow(['日期','净值','日增长率'])
            writer.writerows([r['date'],r['nav'],'' if r.get('daily_growth_pct') in (None,'') else str(r['daily_growth_pct'])+'%'] for r in doc['rows'])
        # Older versions are one recovery archive; live raw evidence remains addressable.
        oldfiles=[x for x in old.iterdir() if x.is_file()] if old.exists() else []
        snapshot=central/'整理前旧版净值.zip'
        if oldfiles:
            with zipfile.ZipFile(snapshot,'x',zipfile.ZIP_DEFLATED) as z:
                for x in oldfiles:
                    if x.name not in (jp.name,cp.name): z.write(x,x.name)
            with zipfile.ZipFile(snapshot) as z:
                assert z.testzip() is None
                for x in oldfiles:
                    if x.name not in (jp.name,cp.name): assert hashlib.sha256(z.read(x.name)).hexdigest()==sha(x)
        if old.exists():
            for child in old.iterdir():
                if child.is_dir():
                    target=central/child.name
                    shutil.move(str(child),str(target))
                    migration.append(dict(old_prefix=rel(old/child.name),new_prefix=rel(target)))
        if code=='100055':
            raw=central/'原始数据/全历史分页_2026-10-20'; raw.mkdir(parents=True,exist_ok=True)
            for source in (STATE/code).glob('page_*.json'): shutil.move(str(source),str(raw/source.name))
        old.mkdir(parents=True,exist_ok=True)
        target=old/f'{code}_全部历史净值.xlsx'
        shutil.copy2(ROOT/f['staged'], target.with_suffix('.replace-tmp'))
        for source in oldfiles: source.unlink()
        target.with_suffix('.replace-tmp').replace(target)
        assert list(old.iterdir())==[target]
        verify(doc,target)
        migration.append(dict(old_prefix=f['history'],new_prefix=rel(central),xlsx_path=rel(target)))
    for report in p['pdfs']:
        if 'target' not in report: continue
        target=ROOT/report['target']; target.parent.mkdir(parents=True,exist_ok=True)
        (ROOT/report['source']).rename(target)
        assert sha(target)==report['sha256']
    for report in p['pdfs']:
        if 'duplicate_of' in report:
            assert sha(ROOT/report['duplicate_of'])==report['sha256']
            (ROOT/report['source']).unlink()
    save(ARCHIVE/'目录迁移记录_2026-10-20.json',dict(mappings=migration,pdfs=p['pdfs']))
    refresh()

def refresh():
    idx=load(INDEX); entries=[]
    oldentries={e['code']:e for e in idx['funds']}
    for folder in sorted((ROOT/'基金').glob('*/*')):
        if not folder.is_dir() or not re.match(r'^\d{6}_',folder.name): continue
        code=folder.name[:6]; jp=CENTRAL/code/f'{code}_全部历史净值.json'; cp=jp.with_suffix('.csv')
        combined=folder/f'{code}_交易记录.xlsx'; legacy=folder/f'{code}_全部历史净值.xlsx'
        has_combined_nav=False
        if combined.exists():
            probe=openpyxl.load_workbook(combined, read_only=True, data_only=True)
            has_combined_nav='全部历史净值' in probe.sheetnames
            probe.close()
        xp=combined if has_combined_nav else legacy
        d=load(jp); verify(d,xp)
        with cp.open(encoding='utf-8-sig', newline='') as handle:
            saved=list(csv.reader(handle))
        assert cp.read_bytes().startswith(b'\xef\xbb\xbf')
        assert saved[0]==['日期','净值','日增长率'] and len(saved)==len(d['rows'])+1
        for actual,row in zip(saved[1:],d['rows']):
            assert actual==[row['date'],str(row['nav']),'' if row.get('daily_growth_pct') in (None,'') else str(row['daily_growth_pct'])+'%']
        if xp==combined:
            assert list(folder.glob(code+'_全部历史净值*'))==[]
        else:
            assert list(folder.glob(code+'_全部历史净值*'))==[legacy]
        assert not (folder/'历史净值').exists()
        e=dict(oldentries.get(code,{})); e.update(code=code,name=d['name'],records=d['count'],first_date=d['first_date'],last_date=d['last_date'],source=d.get('source',d.get('source_type')),json_path=rel(jp),csv_path=rel(cp),xlsx_path=rel(xp),json_sha256=sha(jp),csv_sha256=sha(cp),xlsx_sha256=sha(xp))
        entries.append(e)
    now=datetime.now(timezone(timedelta(hours=8))).isoformat(timespec='seconds')
    policy='基金根目录只保留一份交易记录工作簿，历史净值放在同一文件的「全部历史净值」sheet；机器读取的JSON/CSV和公开核验证据集中保存在公共资料/历史净值/维护数据/，不在基金目录生成重复格式或日期快照。'
    idx.update(updated_at=now,storage_policy=policy,funds=entries)
    save(INDEX,idx)
    text='# 全基金历史净值当前入口\n\n'+policy+'\n\nPDF定期报告按报告所属年份归档，其他公告、合同和招募说明书按文件年份归档。更新日期与实际净值日期分开。\n\n|代码|基金|历史条数|首日|末日|完整表格|\n|---|---|---:|---|---|---|\n'
    for e in entries: text+=f"|{e['code']}|{e['name']}|{e['records']}|{e['first_date']}|{e['last_date']}|[Excel](../../{e['xlsx_path']})|\n"
    INDEX.with_suffix('.md').write_text(text,encoding='utf-8')
    for name in ['六基金历史净值当前入口','四基金历史净值当前入口']:
        ptr=ARCHIVE/(name+'.json')
        if not ptr.exists(): continue
        d=load(ptr); previous={e['code']:e for e in d['funds']}; scope=d.get('scope',list(previous))
        d.update(updated_at=now,storage_policy=policy,funds=[dict(previous.get(e['code'],{}),**e) for e in entries if e['code'] in scope])
        # Immutable historical evidence is resolved through the migration record.
        mappings=load(ARCHIVE/'目录迁移记录_2026-10-20.json')['mappings']
        def replace(value):
            if isinstance(value,str):
                for m in mappings:
                    if value.startswith(m['old_prefix']+'/'): return m['new_prefix']+value[len(m['old_prefix']):]
                return value
            if isinstance(value,list): return [replace(v) for v in value]
            if isinstance(value,dict): return {k:replace(v) for k,v in value.items()}
            return value
        # Keep Excel links under fund directories, unlike machine-data paths.
        d=replace(d)
        for e in d['funds']:
            truth=next(x for x in entries if x['code']==e['code']); e.update(truth)
        save(ptr,d)
        ptr.with_suffix('.md').write_text('# '+name+'\n\n'+policy+'\n\n[全部基金表格入口](全基金历史净值当前入口.md)\n\n'+'\n'.join(f"- {e['code']} {e['name']}：截至{e['last_date']}，[Excel](../../{e['xlsx_path']})。" for e in d['funds'])+'\n',encoding='utf-8')
    print(json.dumps(dict(verified_funds=len(entries),complete_workbooks=len(entries),public_nav_rows=sum(e['records'] for e in entries))))

if __name__=='__main__':
    {'--plan':plan,'--apply':apply,'--refresh':refresh}[sys.argv[1]]()
